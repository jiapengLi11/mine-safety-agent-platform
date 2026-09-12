from __future__ import annotations

import re
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "矿区智能安全监控与处置平台系统设计与开发手册_v1.0.md"
OUTPUT = ROOT / "docs" / "矿区智能安全监控与处置平台系统设计与开发手册_v1.0.docx"
ASSETS = ROOT / "docs" / "assets"
EXISTING_PLATFORM = ASSETS / "label-review-platform.png"

NAVY = "17324D"
TEAL = "168B84"
PALE_BLUE = "EAF2F7"
PALE_TEAL = "E8F5F2"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "D9D9D9"
TEXT = "1D2730"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID_GRAY, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.2

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for name, size, before, after in (
        ("Heading 1", 17, 18, 8),
        ("Heading 2", 13, 13, 6),
        ("Heading 3", 11, 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "MineGuard Agent  矿区智能安全监控与处置平台"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.runs[0]
    set_east_asia_font(header_run, "Microsoft YaHei")
    header_run.font.size = Pt(8.5)
    header_run.font.color.rgb = RGBColor.from_string("5E6C78")
    add_page_number(section.footer.paragraphs[0])


def configure_matplotlib() -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS"]
    plt.rcParams["axes.unicode_minus"] = False


def create_metrics_chart(path: Path) -> None:
    configure_matplotlib()
    labels = ["Precision", "Recall", "mAP50", "mAP50-95"]
    baseline = [83.31, 72.12, 79.64, 55.79]
    reviewed = [84.10, 81.44, 88.34, 63.47]
    x = range(len(labels))
    fig, ax = plt.subplots(figsize=(10.8, 4.4), dpi=180)
    width = 0.34
    bars1 = ax.bar([v - width / 2 for v in x], baseline, width, label="Baseline", color="#6B7F90")
    bars2 = ax.bar([v + width / 2 for v in x], reviewed, width, label="Reviewed v3", color="#168B84")
    ax.set_ylim(0, 100)
    ax.set_ylabel("Percent")
    ax.set_title("Independent test comparison at 640 input size", fontsize=14, weight="bold")
    ax.set_xticks(list(x), labels)
    ax.grid(axis="y", alpha=0.2)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, ncol=2, loc="upper left")
    for bars in (bars1, bars2):
        for bar in bars:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 1.1,
                f"{bar.get_height():.2f}%",
                ha="center",
                va="bottom",
                fontsize=9,
            )
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def create_architecture_diagram(path: Path) -> None:
    configure_matplotlib()
    fig, ax = plt.subplots(figsize=(12, 6.5), dpi=180)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 7)
    ax.axis("off")

    def box(x, y, w, h, title, subtitle, color, text_color="white"):
        patch = FancyBboxPatch(
            (x, y), w, h,
            boxstyle="round,pad=0.03,rounding_size=0.08",
            linewidth=1.2,
            edgecolor=color,
            facecolor=color,
        )
        ax.add_patch(patch)
        ax.text(x + w / 2, y + h * 0.62, title, ha="center", va="center", color=text_color, fontsize=11, weight="bold")
        ax.text(x + w / 2, y + h * 0.30, subtitle, ha="center", va="center", color=text_color, fontsize=8.5)
        return patch

    def arrow(x1, y1, x2, y2, label=""):
        arr = FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=12, linewidth=1.3, color="#506273")
        ax.add_patch(arr)
        if label:
            ax.text((x1 + x2) / 2, (y1 + y2) / 2 + 0.18, label, ha="center", fontsize=8, color="#506273")

    box(0.3, 4.9, 1.8, 1.0, "视频源", "RTSP 摄像头 录像", "#34495E")
    box(2.7, 4.9, 2.0, 1.0, "视觉推理", "Python YOLO26", "#1F6F8B")
    box(5.3, 4.9, 1.4, 1.0, "Kafka", "检测事件", "#C77700")
    box(7.3, 4.55, 2.6, 1.7, "Safety Core", "Spring Boot 模块化单体\n规则 告警 工单 Agent", "#17324D")
    box(10.4, 4.9, 1.3, 1.0, "控制台", "Vue3", "#168B84")
    arrow(2.1, 5.4, 2.7, 5.4)
    arrow(4.7, 5.4, 5.3, 5.4, "frame event")
    arrow(6.7, 5.4, 7.3, 5.4)
    arrow(9.9, 5.4, 10.4, 5.4, "SSE")

    infra = [
        (0.8, 2.35, "MySQL", "业务事实 审计", "#5A7184"),
        (3.0, 2.35, "Redis", "时间窗 幂等", "#B7483B"),
        (5.2, 2.35, "Elasticsearch", "混合检索", "#98752B"),
        (7.8, 2.35, "MinIO", "截图 视频", "#8A496B"),
        (10.0, 2.35, "Observability", "Trace Metrics", "#4B6F44"),
    ]
    for x, y, title, subtitle, color in infra:
        box(x, y, 1.6, 0.9, title, subtitle, color)
        arrow(8.6, 4.55, x + 0.8, y + 0.9)

    box(4.45, 0.45, 3.1, 1.0, "数据治理闭环", "补标平台 数据集 模型评测", "#435C6C")
    arrow(8.0, 4.55, 6.5, 1.45, "feedback")
    arrow(5.5, 1.45, 3.8, 4.9, "released model")
    ax.set_title("MineGuard Agent logical architecture", fontsize=16, weight="bold", pad=12)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_east_asia_font(run, "Microsoft YaHei")
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string("536471")


def add_figure(doc: Document, path: Path, width: float, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_inline_runs(paragraph, text: str) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_east_asia_font(run, "Microsoft YaHei")
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_east_asia_font(run, "Consolas")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string("165D73")
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_east_asia_font(run, "Microsoft YaHei")
            run.bold = True
        else:
            label = token[1:token.index("]")]
            run = paragraph.add_run(label)
            set_east_asia_font(run, "Microsoft YaHei")
            run.font.color.rgb = RGBColor.from_string("176B87")
            run.underline = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_east_asia_font(run, "Microsoft YaHei")


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_east_asia_font(run, "Consolas")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string("23313D")
        if index < len(lines) - 1:
            run.add_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for r_index, values in enumerate(rows):
        row = table.rows[r_index]
        prevent_row_split(row)
        if r_index == 0:
            set_repeat_table_header(row)
        for c_index, value in enumerate(values):
            cell = row.cells[c_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margin(cell)
            if r_index == 0:
                set_cell_shading(cell, NAVY)
            elif r_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_east_asia_font(run, "Microsoft YaHei")
            run.font.size = Pt(8.5)
            if r_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw.append(lines[index].strip())
        index += 1
    rows = []
    for position, line in enumerate(raw):
        values = [part.strip() for part in line.strip("|").split("|")]
        if position == 1 and all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            continue
        rows.append(values)
    return rows, index


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(54)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("矿区智能安全监控与处置平台\n系统设计与开发手册")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(12)
    subtitle.paragraph_format.space_after = Pt(32)
    run = subtitle.add_run("MineGuard Agent 版本 1.0")
    set_east_asia_font(run, "Microsoft YaHei")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    run.font.bold = True

    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(22)
    summary_run = summary.add_run(
        "用于指导视觉检测、时序告警、RAG 证据检索、受控 Agent、工单闭环和模型反馈系统的开发、测试与演示。"
    )
    set_east_asia_font(summary_run, "Microsoft YaHei")
    summary_run.font.size = Pt(12)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    values = [
        ("设计基线", "视觉负责感知 规则负责判断 Agent 负责受控编排"),
        ("当前基础", "29,071 张图像 174,426 次推理 30,183 条审核任务"),
        ("技术主线", "Java 21 Spring Boot Vue3 Kafka Redis MySQL Spring AI"),
        ("编制日期", "2026 年 9 月 12 日"),
    ]
    for index, (left, right) in enumerate(values):
        for c_index, value in enumerate((left, right)):
            cell = table.rows[index].cells[c_index]
            set_cell_border(cell, "D7DEE3")
            set_cell_margin(cell, top=130, bottom=130)
            set_cell_shading(cell, PALE_TEAL if c_index == 0 else "FFFFFF")
            run = cell.paragraphs[0].add_run(value)
            set_east_asia_font(run, "Microsoft YaHei")
            run.font.size = Pt(9.5 if c_index == 0 else 10)
            run.bold = c_index == 0
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_navigation(doc: Document, headings: list[str]) -> None:
    doc.add_heading("内容导航", level=1)
    paragraph = doc.add_paragraph(
        "本手册从产品边界开始，依次给出架构、数据、业务状态、AI 安全、测试和开发任务。开发时以第二十四章和第二十五章作为迭代入口。"
    )
    paragraph.paragraph_format.space_after = Pt(10)
    rows = []
    half = (len(headings) + 1) // 2
    for index in range(half):
        left = headings[index]
        right = headings[index + half] if index + half < len(headings) else ""
        rows.append([left, right])
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_index, values in enumerate(rows):
        for c_index, value in enumerate(values):
            cell = table.rows[r_index].cells[c_index]
            set_cell_border(cell, "E2E7EA", "4")
            set_cell_margin(cell, top=75, bottom=75)
            if r_index % 2 == 1:
                set_cell_shading(cell, "F7F9FA")
            run = cell.paragraphs[0].add_run(value)
            set_east_asia_font(run, "Microsoft YaHei")
            run.font.size = Pt(8.5)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    metrics = ASSETS / "model_metrics_comparison.png"
    architecture = ASSETS / "logical_architecture.png"
    create_metrics_chart(metrics)
    create_architecture_diagram(architecture)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = [line[3:].strip() for line in lines if line.startswith("## ")]

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_navigation(doc, headings)

    start = next(index for index, line in enumerate(lines) if line.startswith("## 1 "))
    index = start
    in_code = False
    code_lines: list[str] = []
    page_break_sections = {"5", "8", "13", "17", "20", "24"}

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if stripped == "[[FIGURE_METRICS]]":
            add_figure(doc, metrics, 6.65, "图 1  审核数据重训模型与基线模型独立测试对比")
            index += 1
            continue
        if stripped == "[[FIGURE_PLATFORM]]":
            if EXISTING_PLATFORM.exists():
                add_figure(doc, EXISTING_PLATFORM, 6.65, "图 2  已运行的多人补标审核平台进度看板")
            index += 1
            continue
        if stripped == "[[FIGURE_ARCHITECTURE]]":
            add_figure(doc, architecture, 6.7, "图 3  MineGuard Agent 逻辑架构")
            index += 1
            continue

        if raw.startswith("## "):
            text = raw[3:].strip()
            section_match = re.match(r"(\d+)\s", text)
            heading = doc.add_heading(text, level=1)
            if section_match and section_match.group(1) in page_break_sections:
                heading.paragraph_format.page_break_before = True
            index += 1
            continue
        if raw.startswith("### "):
            doc.add_heading(raw[4:].strip(), level=2)
            index += 1
            continue
        if raw.startswith("#### "):
            doc.add_heading(raw[5:].strip(), level=3)
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        if re.match(r"^-\s+", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline_runs(paragraph, re.sub(r"^-\s+", "", stripped))
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline_runs(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(stripped.lstrip("> "))
            set_east_asia_font(run, "Microsoft YaHei")
            run.italic = True
            run.font.color.rgb = RGBColor.from_string("445A66")
            index += 1
            continue
        if stripped:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = False
            add_inline_runs(paragraph, stripped)
        index += 1

    core = doc.core_properties
    core.title = "矿区智能安全监控与处置平台系统设计与开发手册"
    core.subject = "MineGuard Agent 系统设计与开发基线"
    core.author = "李嘉鹏"
    core.keywords = "Spring Boot, YOLO26, RAG, Agent, Kafka, Redis, Mine Safety"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
